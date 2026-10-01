"""Stable report-designer scope, disclosure, and column contracts for #565."""

import csv
import io
from contextlib import ExitStack, contextmanager
from datetime import datetime
from datetime import timezone as dt_timezone
from types import SimpleNamespace
from unittest.mock import MagicMock, Mock, call, patch

from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.test import SimpleTestCase

from assets.reports import AssetDepreciationReportProvider
from core.context import set_current_all_accessible, set_current_tenant, set_current_tenant_group
from core.reports.contracts import (
    ReportDefinition,
    ReportPermissionDenied,
    ReportRequest,
    ReportResult,
    record_window_state,
)
from core.reports.exporters import report_xlsx_bytes
from core.reports.orchestration import (
    REPORT_COMPILATION_OPERATION,
    _enforce_report_provider_permissions,
    _global_report_permissions_cover,
    _live_tenant_ids,
    _principal_covers_report_permissions,
    _report_compilation_authorized,
    _resolve_report_scope,
    build_report_context,
)
from core.reports.registry import get_report_provider
from core.reports.rendering import render_report_csv, render_report_html
from extras.forms import ReportTemplateForm
from extras.tasks.reports import _attachment_email_body
from extras.views import ReportTemplateDownloadView, _apply_report_disclosure_headers
from inventory.reports import HardwareInventoryReportProvider
from procurement.reports import ContractRenewalsReportProvider
from subscriptions.reports import SubscriptionRenewalsReportProvider


class ReportCompileScopeAuthorizationTests(SimpleTestCase):
    def setUp(self):
        self.tenant_a = SimpleNamespace(pk=1, id=1)
        self.tenant_b = SimpleNamespace(pk=2, id=2)
        self.tenant_c = SimpleNamespace(pk=3, id=3)
        self.user = Mock(pk=41, is_superuser=False)

    def test_single_pinned_tenant_is_allowed_when_reachable(self):
        with (
            patch("core.reports.orchestration.get_current_user", return_value=self.user),
            patch("core.reports.orchestration.accessible_tenant_ids", return_value={self.tenant_b.pk}),
            patch("core.reports.orchestration._tenant_scope_reach_is_valid") as cross_tenant_check,
        ):
            scope = _resolve_report_scope(self.tenant_a, [self.tenant_b])

        self.assertEqual(scope, [self.tenant_b])
        cross_tenant_check.assert_not_called()

    def test_single_foreign_tenant_permission_is_checked_inside_its_task_context(self):
        self.user.has_perm.return_value = True
        task_context = patch("core.tasks.context.TaskContext")
        with (
            patch("core.reports.orchestration.get_current_user", return_value=self.user),
            patch("core.reports.orchestration.accessible_tenant_ids", return_value=set()),
            patch("itambox.middleware.get_current_user", return_value=self.user),
            task_context as task_context_class,
        ):
            self.assertEqual(_resolve_report_scope(None, [self.tenant_b]), [self.tenant_b])

        task_context_class.assert_called_once_with(
            tenant_id=self.tenant_b.pk,
            user_id=self.user.pk,
            operation="reports.scope",
        )
        self.user.has_perm.assert_called_once_with("reports.view_cross_tenant_reports", obj=self.tenant_b)

    def test_multi_tenant_pinned_scope_requires_permission_for_every_tenant(self):
        self.user.has_perm.side_effect = lambda permission, obj=None: obj is self.tenant_a
        with (
            patch("core.reports.orchestration.get_current_user", return_value=self.user),
            patch("core.reports.orchestration.accessible_tenant_ids", return_value={1, 2, 3}),
            patch("itambox.middleware.get_current_user", return_value=self.user),
            patch("core.tasks.context.TaskContext") as task_context,
        ):
            with self.assertRaisesRegex(ReportPermissionDenied, "reports.view_cross_tenant_reports"):
                _resolve_report_scope(self.tenant_a, [self.tenant_a, self.tenant_b])

        self.assertEqual(
            task_context.call_args_list,
            [
                call(tenant_id=1, user_id=self.user.pk, operation="reports.scope"),
                call(tenant_id=2, user_id=self.user.pk, operation="reports.scope"),
            ],
        )
        self.assertEqual(
            self.user.has_perm.call_args_list,
            [
                call("reports.view_cross_tenant_reports", obj=self.tenant_a),
                call("reports.view_cross_tenant_reports", obj=self.tenant_b),
            ],
        )

        self.user.has_perm.reset_mock()
        self.user.has_perm.side_effect = lambda permission, obj=None: obj in (self.tenant_a, self.tenant_b)
        with (
            patch("core.reports.orchestration.get_current_user", return_value=self.user),
            patch("core.reports.orchestration.accessible_tenant_ids", return_value=set()),
            patch("itambox.middleware.get_current_user", return_value=self.user),
            patch("core.tasks.context.TaskContext") as task_context,
        ):
            self.assertEqual(
                _resolve_report_scope(self.tenant_a, [self.tenant_a, self.tenant_b]),
                [self.tenant_a, self.tenant_b],
            )
        self.assertEqual(task_context.call_count, 2)

    def test_system_context_allows_only_one_pinned_tenant_matching_active_scope(self):
        with patch("core.reports.orchestration.get_current_user", return_value=None):
            self.assertEqual(_resolve_report_scope(None, [self.tenant_b]), [self.tenant_b])
            self.assertEqual(_resolve_report_scope(self.tenant_b, [self.tenant_b]), [self.tenant_b])
            for active_tenant, pinned in (
                (self.tenant_a, [self.tenant_b]),
                (None, [self.tenant_a, self.tenant_b]),
            ):
                with self.subTest(active_tenant=active_tenant, pinned=pinned):
                    with self.assertRaises(ReportPermissionDenied):
                        _resolve_report_scope(active_tenant, pinned)

    def test_unresolvable_pinned_tenant_entry_is_refused(self):
        with patch("core.reports.orchestration.get_current_user", return_value=self.user):
            with self.assertRaises(ReportPermissionDenied):
                _resolve_report_scope(None, [SimpleNamespace(pk=None)])

    def test_empty_scope_is_global_only_for_permission_holders(self):
        self.user.has_perm.return_value = True
        with patch("core.reports.orchestration.get_current_user", return_value=self.user):
            self.assertEqual(_resolve_report_scope(self.tenant_a, []), [])
        self.user.has_perm.return_value = False
        with patch("core.reports.orchestration.get_current_user", return_value=self.user):
            self.assertEqual(_resolve_report_scope(self.tenant_a, []), [self.tenant_a])
            with self.assertRaises(PermissionError):
                _resolve_report_scope(None, [])


class _PinnedDownloadProvider(ReportDefinition):
    report_type = "asset_summary"
    default_columns = ("asset_tag",)
    cells = {"asset_tag": lambda record, request: record}

    def __init__(self, specification_export=None, *, truncated=False, total_rows=None):
        self.specification_export = specification_export
        self.truncated = truncated
        self.total_rows = total_rows
        self.request = None

    def build(self, request):
        self.request = request
        rows = [{"Asset Tag": f"tenant-{tenant.pk}"} for tenant in request.filter_tenants] or [
            {"Asset Tag": "global-row"}
        ]
        return ReportResult(
            rows=rows,
            truncated=self.truncated,
            total_rows=self.total_rows,
            specification_export=self.specification_export,
        )


def _download_template(pinned_ids):
    return SimpleNamespace(
        pk=17,
        name="Pinned Report",
        description="",
        report_type="asset_summary",
        included_columns=["asset_tag"],
        group_by_field="",
        style_preset="default",
        template_content="",
        persisted_filter_tenant_ids=Mock(return_value=pinned_ids),
    )


class ReportPinnedDownloadTests(SimpleTestCase):
    def setUp(self):
        self.ambient_tenant = SimpleNamespace(pk=1, id=1)
        self.tenant_b = SimpleNamespace(pk=2, id=2)
        self.tenant_c = SimpleNamespace(pk=3, id=3)
        self.template = _download_template([self.tenant_b.pk, self.tenant_c.pk])
        self.user = Mock(pk=41, is_superuser=False)
        self.tenant_queryset = Mock()
        self.tenant_queryset.order_by.return_value = [self.tenant_b, self.tenant_c]
        self.tenant_manager = Mock()
        self.tenant_manager.filter.return_value = self.tenant_queryset
        self.tenant_model = SimpleNamespace(_base_manager=self.tenant_manager)

    def _view_patches(self, provider, *, authorized):
        return (
            patch("extras.views.get_object_or_404", return_value=self.template),
            patch("extras.views.get_current_tenant", return_value=self.ambient_tenant),
            patch("django.apps.apps.get_model", return_value=self.tenant_model),
            patch("core.reports.orchestration.get_current_user", return_value=self.user),
            patch("core.reports.orchestration.accessible_tenant_ids", return_value=set()),
            patch("core.reports.orchestration._tenant_scope_reach_is_valid", return_value=authorized),
            # Domain permission enforcement has separate coverage; these cases verify the pinned scope machinery.
            patch("core.reports.orchestration._enforce_report_provider_permissions"),
            patch("core.reports.orchestration.get_report_provider", return_value=provider),
        )

    @contextmanager
    def _patched_download(self, provider, *, authorized):
        with ExitStack() as stack:
            for patcher in self._view_patches(provider, authorized=authorized):
                stack.enter_context(patcher)
            yield

    def test_foreign_pinned_download_returns_403_without_truncating_the_constellation(self):
        provider = _PinnedDownloadProvider()
        with self._patched_download(provider, authorized=False):
            response = ReportTemplateDownloadView().get(SimpleNamespace(GET={"format": "csv"}), self.template.pk)

        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.content, b"You may not view this report's data.")
        self.template.persisted_filter_tenant_ids.assert_called_once_with()
        self.tenant_manager.filter.assert_called_once_with(
            pk__in=[self.tenant_b.pk, self.tenant_c.pk],
            deleted_at__isnull=True,
        )
        self.assertIsNone(provider.request)

    def test_authorized_download_compiles_every_persisted_tenant(self):
        provider = _PinnedDownloadProvider()
        with self._patched_download(provider, authorized=True):
            response = ReportTemplateDownloadView().get(SimpleNamespace(GET={"format": "csv"}), self.template.pk)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(provider.request.filter_tenants, (self.tenant_b, self.tenant_c))
        self.assertIn(b"tenant-2", response.content)
        self.assertIn(b"tenant-3", response.content)

    def test_partially_deleted_constellation_fails_closed(self):
        provider = _PinnedDownloadProvider()
        # B survives the unscoped read; C was soft-deleted and drops out.
        self.tenant_queryset.order_by.return_value = [self.tenant_b]
        with self.assertLogs("extras.views", level="ERROR") as logs:
            with self._patched_download(provider, authorized=True):
                response = ReportTemplateDownloadView().get(SimpleNamespace(GET={"format": "csv"}), self.template.pk)

        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.content, b"You may not view this report's data.")
        self.assertIn("soft-deleted; refusing compilation", "\n".join(logs.output))
        self.assertIsNone(provider.request)

    def test_machine_csv_requires_an_export_and_csv_keeps_the_frozen_machine_export_path(self):
        absent_export_provider = _PinnedDownloadProvider(truncated=True, total_rows=600)
        template = _download_template([])
        self.template = template
        with self._patched_download(absent_export_provider, authorized=True):
            response = ReportTemplateDownloadView().get(SimpleNamespace(GET={"format": "machine_csv"}), template.pk)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.content, b"This report does not provide a machine-format export.")

        from assets.services.specification_consumers.exporting import MachineExportResult, machine_csv_bytes

        machine_export = MachineExportResult(
            columns=("asset.spec.value",),
            rows=({"asset.spec.value": "24.000"},),
            metadata=(),
        )
        export_provider = _PinnedDownloadProvider(
            machine_export,
            truncated=True,
            total_rows=600,
        )
        with self._patched_download(export_provider, authorized=True):
            response = ReportTemplateDownloadView().get(SimpleNamespace(GET={"format": "csv"}), template.pk)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, machine_csv_bytes(machine_export))
        self.assertEqual(response["X-Report-Truncated"], "true")
        self.assertEqual(response["X-Report-Row-Window"], "500")
        self.assertEqual(response["X-Report-Total-Rows"], "600")


class _WindowQueryset:
    def __init__(self, total, row_count):
        self.total = total
        self.records = [object() for _ in range(row_count)]

    def __getitem__(self, item):
        return self.records[item]

    def count(self):
        return self.total


class ReportDisclosureTests(SimpleTestCase):
    def test_full_row_windows_always_have_a_determinate_total(self):
        queryset = _WindowQueryset(total=4, row_count=3)
        self.assertEqual(record_window_state(queryset, queryset[:3], 3), (True, 4))
        self.assertEqual(record_window_state(queryset, queryset[:2], 3), (False, None))
        exact_queryset = _WindowQueryset(total=3, row_count=3)
        self.assertEqual(record_window_state(exact_queryset, exact_queryset[:3], 3), (False, 3))

    def test_provider_overrides_carry_determinate_truncation_totals(self):
        request_template = SimpleNamespace(
            include_summary_cards=False,
            include_distribution_chart=False,
            group_by_field="",
        )
        providers = (
            AssetDepreciationReportProvider(),
            SubscriptionRenewalsReportProvider(),
            ContractRenewalsReportProvider(),
        )
        for provider in providers:
            provider.row_limit = 1
            queryset = _WindowQueryset(total=7, row_count=3)
            request = ReportRequest(
                template=request_template,
                active_tenant=None,
                filter_tenants=(),
                columns=provider.default_columns,
                user=None,
                as_of=datetime(2026, 1, 1),
            )
            with (
                patch.object(provider, "get_queryset", return_value=queryset),
                patch.object(provider, "build_rows", return_value=[{"Asset Tag": "one"}]),
            ):
                result = provider.build(request)
            self.assertTrue(result.truncated, provider.report_type)
            self.assertEqual(result.total_rows, 7, provider.report_type)

        hardware = HardwareInventoryReportProvider()
        hardware.row_limit = 1
        catalogues = (
            ("Accessory", _WindowQueryset(total=4, row_count=2)),
            ("Component", _WindowQueryset(total=3, row_count=2)),
        )
        request = ReportRequest(
            template=request_template,
            active_tenant=None,
            filter_tenants=(),
            columns=hardware.default_columns,
            user=None,
            as_of=datetime(2026, 1, 1),
        )
        with (
            patch.object(hardware, "get_queryset", return_value=catalogues),
            patch.object(hardware, "_catalogue_records", return_value=[object()]),
            patch.object(hardware, "build_rows", return_value=[{"Item": "one"}]),
            patch.object(hardware, "_summary_cards", return_value=[]),
            patch.object(hardware, "_chart", return_value=""),
        ):
            result = hardware.build(request)
        self.assertTrue(result.truncated)
        self.assertEqual(result.total_rows, 7)

    def test_context_discloses_sample_and_row_window_state(self):
        class Provider(ReportDefinition):
            report_type = "asset_summary"
            default_columns = ("asset_tag",)

            def build(self, request):
                return ReportResult(rows=[{"Asset Tag": "A-1"}], truncated=True, total_rows=750)

        template = SimpleNamespace(
            report_type="asset_summary",
            name="Assets",
            description="",
            included_columns=["asset_tag"],
            group_by_field="",
            style_preset="default",
        )
        with (
            # Domain permission enforcement has separate coverage; this case verifies disclosure assembly.
            patch("core.reports.orchestration._enforce_report_provider_permissions"),
            patch("core.reports.orchestration.get_current_user", return_value=None),
            patch("core.reports.orchestration.get_report_provider", return_value=Provider()),
        ):
            *_, context_data = build_report_context(template, active_tenant=SimpleNamespace(pk=1))

        self.assertEqual(context_data["is_sample"], False)
        self.assertEqual(context_data["row_limit"], 500)
        self.assertEqual(context_data["truncated"], True)
        self.assertEqual(context_data["total_rows"], 750)
        self.assertIn("750", context_data["disclosure_text"])

        class SampleProvider(Provider):
            def build(self, request):
                return ReportResult(rows=[{"Asset Tag": "A-1"}], is_sample=True)

        with (
            # Domain permission enforcement has separate coverage; this case verifies disclosure assembly.
            patch("core.reports.orchestration._enforce_report_provider_permissions"),
            patch("core.reports.orchestration.get_current_user", return_value=None),
            patch("core.reports.orchestration.get_report_provider", return_value=SampleProvider()),
        ):
            *_, sample_context = build_report_context(template, active_tenant=SimpleNamespace(pk=1))
        self.assertTrue(sample_context["is_sample"])
        self.assertFalse(sample_context["truncated"])
        self.assertIsNone(sample_context["total_rows"])
        self.assertIn("sample data", sample_context["disclosure_text"])


class _ScopeQuery(list):
    """Queryset stand-in: iterable of tenants, chained ``filter`` returns itself."""

    def filter(self, *args, **kwargs):
        return self


class _UnreadableIdentity:
    """Principal stand-in whose ``pk`` cannot be resolved for logging."""

    is_authenticated = True
    is_active = True
    is_superuser = False

    @property
    def pk(self):
        raise RuntimeError("unreadable pk")


class _UnreadableTenant:
    """Tenant stand-in whose ``pk`` cannot be resolved for logging."""

    @property
    def pk(self):
        raise RuntimeError("unreadable pk")


class _UnreadableReportProvider:
    """Provider stand-in whose ``report_type`` cannot be resolved for logging."""

    @property
    def report_type(self):
        raise RuntimeError("unreadable report type")

    def required_permissions(self):
        return ("assets.view_asset",)


class ReportProviderDomainPermissionTests(SimpleTestCase):
    def setUp(self):
        self.tenant_a = SimpleNamespace(pk=1)
        self.tenant_b = SimpleNamespace(pk=2)
        self.tenant_c = SimpleNamespace(pk=3)
        self.user = Mock(pk=41, is_authenticated=True, is_active=True, is_superuser=False)

    def test_hardware_inventory_requires_every_declared_permission(self):
        provider = get_report_provider("hardware_inventory")
        permissions = provider.required_permissions()
        self.assertEqual(len(permissions), 3)
        held_permissions = set(permissions) - {"inventory.view_component"}
        with (
            patch("core.reports.orchestration.get_current_user", return_value=self.user),
            patch(
                "core.reports.orchestration._principal_covers_report_permissions",
                side_effect=lambda _user, _tenant, declared: all(
                    permission in held_permissions for permission in declared
                ),
            ),
        ):
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [self.tenant_a], self.tenant_a)

        with (
            patch("core.reports.orchestration.get_current_user", return_value=self.user),
            patch("core.reports.orchestration._principal_covers_report_permissions", return_value=True),
        ):
            _enforce_report_provider_permissions(provider, [self.tenant_a], self.tenant_a)

    def test_each_pinned_tenant_is_checked_and_one_missing_permission_denies(self):
        provider = get_report_provider("asset_summary")
        with (
            patch("core.reports.orchestration.get_current_user", return_value=self.user),
            patch(
                "core.reports.orchestration._principal_covers_report_permissions",
                side_effect=[True, False],
            ) as check,
        ):
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [self.tenant_a, self.tenant_b], self.tenant_c)

        self.assertEqual([item.args[1] for item in check.call_args_list], [self.tenant_a, self.tenant_b])

    def test_per_tenant_resolver_failures_are_denials(self):
        with patch("core.tasks.context.TaskContext", side_effect=PermissionDenied):
            self.assertFalse(_principal_covers_report_permissions(self.user, self.tenant_a, ("assets.view_asset",)))

        bound_user = Mock(pk=self.user.pk, is_authenticated=True, is_active=True)
        task_context = MagicMock()
        task_context.return_value.__enter__.return_value = Mock()
        with (
            patch("core.tasks.context.TaskContext", task_context),
            patch("core.reports.orchestration.get_current_user", return_value=bound_user),
        ):
            self.assertTrue(
                _principal_covers_report_permissions(
                    self.user,
                    self.tenant_a,
                    ("assets.view_asset", "assets.view_asset"),
                )
            )
        task_context.assert_called_once_with(tenant_id=1, user_id=self.user.pk, operation="reports.permissions")
        self.assertEqual(bound_user.has_perm.call_count, 2)
        self.assertTrue(all(item.kwargs["obj"] is self.tenant_a for item in bound_user.has_perm.call_args_list))

    def test_active_superuser_passes_and_inactive_superuser_fails(self):
        provider = get_report_provider("asset_summary")
        active_superuser = SimpleNamespace(pk=7, is_authenticated=True, is_active=True, is_superuser=True)
        with (
            patch("core.reports.orchestration.get_current_user", return_value=active_superuser),
            patch("core.reports.orchestration._principal_covers_report_permissions") as check,
        ):
            _enforce_report_provider_permissions(provider, [], self.tenant_a)
        check.assert_not_called()

        inactive_superuser = SimpleNamespace(pk=8, is_authenticated=True, is_active=False, is_superuser=True)
        with patch("core.reports.orchestration.get_current_user", return_value=inactive_superuser):
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [], self.tenant_a)

    def test_actorless_single_tenant_requires_each_exact_system_authorization(self):
        provider = get_report_provider("hardware_inventory")
        permissions = provider.required_permissions()
        with (
            patch("core.reports.orchestration.get_current_user", return_value=None),
            patch("core.reports.orchestration.has_valid_system_authorization", return_value=False),
        ):
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [], self.tenant_a)

        with (
            patch("core.reports.orchestration.get_current_user", return_value=None),
            patch("core.reports.orchestration.has_valid_system_authorization", return_value=True) as check,
        ):
            _enforce_report_provider_permissions(provider, [], self.tenant_a)
        self.assertEqual(
            check.call_args_list,
            [
                call(
                    tenant_id=self.tenant_a.pk,
                    permission=permission,
                    operation=REPORT_COMPILATION_OPERATION,
                )
                for permission in permissions
            ],
        )

        # An unresolvable actorless tenant identity never authorizes.
        with (
            patch("core.reports.orchestration.get_current_user", return_value=None),
            patch("core.reports.orchestration.has_valid_system_authorization", return_value=True),
        ):
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [], _UnreadableTenant())
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [], SimpleNamespace(pk=None))

        # A crashing system-authorization validator denies.
        with (
            patch("core.reports.orchestration.get_current_user", return_value=None),
            patch(
                "core.reports.orchestration.has_valid_system_authorization",
                side_effect=RuntimeError("validator failed"),
            ),
        ):
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [], self.tenant_a)

    def test_actorless_multi_tenant_and_global_scopes_are_denied(self):
        provider = get_report_provider("asset_summary")
        with (
            patch("core.reports.orchestration.get_current_user", return_value=None),
            patch("core.reports.orchestration.has_valid_system_authorization", return_value=True),
        ):
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [self.tenant_a, self.tenant_b], self.tenant_a)
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [], None)

    def test_global_aggregation_requires_all_live_tenants_and_unexpired_grants(self):
        tenant_query = Mock()
        tenant_query.values_list.return_value = (1, 2)
        tenant_model_value = SimpleNamespace(_base_manager=SimpleNamespace(filter=Mock(return_value=tenant_query)))
        permissions = ("assets.view_asset",)
        with (
            patch("core.reports.orchestration.tenant_model", return_value=tenant_model_value),
            patch(
                "core.reports.orchestration.build_accessible_tenant_permissions_map",
                return_value={1: (frozenset(permissions), None), 2: (frozenset(permissions), None)},
            ),
        ):
            self.assertTrue(_global_report_permissions_cover(self.user, permissions))

        with (
            patch("core.reports.orchestration.tenant_model", return_value=tenant_model_value),
            patch(
                "core.reports.orchestration.build_accessible_tenant_permissions_map",
                return_value={1: (frozenset(permissions), None), 2: (frozenset(), None)},
            ),
        ):
            self.assertFalse(_global_report_permissions_cover(self.user, permissions))

        expiry = datetime(2025, 1, 1, tzinfo=dt_timezone.utc)
        with (
            patch("core.reports.orchestration.tenant_model", return_value=tenant_model_value),
            patch("core.reports.orchestration.timezone.now", return_value=datetime(2026, 1, 1, tzinfo=dt_timezone.utc)),
            patch(
                "core.reports.orchestration.build_accessible_tenant_permissions_map",
                return_value={1: (frozenset(permissions), None), 2: (frozenset(permissions), expiry)},
            ),
        ):
            self.assertFalse(_global_report_permissions_cover(self.user, permissions))

        tenant_query.values_list.return_value = ()
        with patch("core.reports.orchestration.tenant_model", return_value=tenant_model_value):
            self.assertFalse(_global_report_permissions_cover(self.user, permissions))

        tenant_query.values_list.side_effect = RuntimeError("tenant enumeration failed")
        with patch("core.reports.orchestration.tenant_model", return_value=tenant_model_value):
            self.assertFalse(_global_report_permissions_cover(self.user, permissions))

        tenant_query.values_list.side_effect = None
        tenant_query.values_list.return_value = (1, 2)
        with (
            patch("core.reports.orchestration.tenant_model", return_value=tenant_model_value),
            patch(
                "core.reports.orchestration.build_accessible_tenant_permissions_map",
                side_effect=RuntimeError("permission map failed"),
            ),
        ):
            self.assertFalse(_global_report_permissions_cover(self.user, permissions))

        with (
            patch("core.reports.orchestration.tenant_model", return_value=tenant_model_value),
            patch(
                "core.reports.orchestration.build_accessible_tenant_permissions_map",
                return_value=["not", "a", "mapping"],
            ),
        ):
            self.assertFalse(_global_report_permissions_cover(self.user, permissions))

        with (
            patch("core.reports.orchestration.tenant_model", return_value=tenant_model_value),
            patch(
                "core.reports.orchestration.build_accessible_tenant_permissions_map",
                return_value={1: "unshaped", 2: (frozenset(permissions), None)},
            ),
        ):
            self.assertFalse(_global_report_permissions_cover(self.user, permissions))

    def test_unscoped_non_superuser_scope_is_unrepresentable_and_denies(self):
        """A non-superuser without a resolved scope reads no rows; never widen to live tenants."""
        provider = get_report_provider("asset_summary")
        tenant_model_value = Mock()
        with (
            patch("core.reports.orchestration.get_current_user", return_value=self.user),
            patch("core.reports.orchestration.tenant_model", tenant_model_value),
        ):
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [], None)
        tenant_model_value.assert_not_called()

        active_superuser = SimpleNamespace(pk=7, is_authenticated=True, is_active=True, is_superuser=True)
        with (
            patch("core.reports.orchestration.get_current_user", return_value=active_superuser),
            patch("core.reports.orchestration.tenant_model", tenant_model_value),
        ):
            _enforce_report_provider_permissions(provider, [], None)
        tenant_model_value.assert_not_called()

    def test_all_accessible_scope_checks_only_the_reachable_tenants(self):
        """The "all accessible tenants" scope resolves like the managers: A+B, never a foreign C."""
        provider = get_report_provider("asset_summary")
        tenant_query = _ScopeQuery([self.tenant_a, self.tenant_b])
        tenant_model_value = SimpleNamespace(_base_manager=SimpleNamespace(filter=Mock(return_value=tenant_query)))
        self.addCleanup(set_current_all_accessible, False)
        set_current_all_accessible(True)
        with (
            patch("core.reports.orchestration.get_current_user", return_value=self.user),
            patch("core.reports.orchestration.tenant_model", return_value=tenant_model_value),
            patch("core.reports.orchestration.accessible_tenant_ids", return_value={1, 2}),
            patch(
                "core.reports.orchestration._principal_covers_report_permissions",
                return_value=True,
            ) as check,
        ):
            _enforce_report_provider_permissions(provider, [], None)
        self.assertEqual([item.args[1] for item in check.call_args_list], [self.tenant_a, self.tenant_b])

        with (
            patch("core.reports.orchestration.get_current_user", return_value=self.user),
            patch("core.reports.orchestration.tenant_model", return_value=tenant_model_value),
            patch("core.reports.orchestration.accessible_tenant_ids", return_value={1, 2}),
            patch(
                "core.reports.orchestration._principal_covers_report_permissions",
                side_effect=[True, False],
            ),
        ):
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [], None)

        with (
            patch("core.reports.orchestration.get_current_user", return_value=self.user),
            patch("core.reports.orchestration.tenant_model", return_value=tenant_model_value),
            patch("core.reports.orchestration.accessible_tenant_ids", return_value={1, 2}),
            patch(
                "core.reports.orchestration._principal_covers_report_permissions",
                side_effect=RuntimeError("resolver crash"),
            ),
        ):
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [], None)

    def test_tenant_group_scope_checks_only_the_reachable_subtree_tenants(self):
        """A tenant-group scope resolves to the accessible tenants of its live subtree."""
        provider = get_report_provider("asset_summary")
        group = SimpleNamespace(pk=7)
        tenant_query = _ScopeQuery([self.tenant_a, self.tenant_b])
        tenant_model_value = SimpleNamespace(_base_manager=SimpleNamespace(filter=Mock(return_value=tenant_query)))
        self.addCleanup(set_current_tenant_group, None)
        set_current_tenant_group(group)
        with (
            patch("core.reports.orchestration.get_current_user", return_value=self.user),
            patch("core.reports.orchestration.tenant_model", return_value=tenant_model_value),
            patch("core.reports.orchestration.accessible_tenant_ids", return_value={1, 2, 3}),
            patch("core.reports.orchestration.get_descendant_tenant_group_ids", return_value={7, 8}) as subtree,
            patch(
                "core.reports.orchestration._principal_covers_report_permissions",
                return_value=True,
            ) as check,
        ):
            _enforce_report_provider_permissions(provider, [], None)
        subtree.assert_called_once_with(7, live_only=True)
        self.assertEqual([item.args[1] for item in check.call_args_list], [self.tenant_a, self.tenant_b])

        with (
            patch("core.reports.orchestration.get_current_user", return_value=self.user),
            patch("core.reports.orchestration.tenant_model", return_value=tenant_model_value),
            patch("core.reports.orchestration.accessible_tenant_ids", return_value={1, 2, 3}),
            patch("core.reports.orchestration.get_descendant_tenant_group_ids", return_value={7, 8}),
            patch(
                "core.reports.orchestration._principal_covers_report_permissions",
                side_effect=[True, False],
            ),
        ):
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [], None)

    def test_principal_coverage_identity_and_bound_resolution_edges(self):
        """Unresolvable identities and bound principals fail closed."""
        task_context = MagicMock()
        task_context.return_value.__enter__.return_value = Mock()
        self.assertFalse(
            _principal_covers_report_permissions(_UnreadableIdentity(), self.tenant_a, ("assets.view_asset",))
        )
        self.assertFalse(_principal_covers_report_permissions(self.user, _UnreadableTenant(), ("assets.view_asset",)))
        self.assertFalse(
            _principal_covers_report_permissions(self.user, SimpleNamespace(pk=None), ("assets.view_asset",))
        )
        self.assertFalse(
            _principal_covers_report_permissions(SimpleNamespace(pk=None), self.tenant_a, ("assets.view_asset",))
        )

        for bound_user in (
            None,
            SimpleNamespace(is_authenticated=False, is_active=True),
            SimpleNamespace(is_authenticated=True, is_active=False),
        ):
            with (
                patch("core.tasks.context.TaskContext", task_context),
                patch("core.reports.orchestration.get_current_user", return_value=bound_user),
            ):
                self.assertFalse(_principal_covers_report_permissions(self.user, self.tenant_a, ("assets.view_asset",)))

        failing_user = Mock(pk=41, is_authenticated=True, is_active=True)
        failing_user.has_perm.side_effect = RuntimeError("permission backend failed")
        with (
            patch("core.tasks.context.TaskContext", task_context),
            patch("core.reports.orchestration.get_current_user", return_value=failing_user),
        ):
            self.assertFalse(_principal_covers_report_permissions(self.user, self.tenant_a, ("assets.view_asset",)))

    def test_denial_logging_tolerates_unreadable_identities(self):
        with (
            patch("core.reports.orchestration.get_current_user", return_value=_UnreadableIdentity()),
            patch("core.reports.orchestration._principal_covers_report_permissions", return_value=False),
        ):
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(_UnreadableReportProvider(), [], _UnreadableTenant())

    def test_conflicting_ambient_scope_states_fail_closed(self):
        provider = get_report_provider("asset_summary")
        self.addCleanup(set_current_tenant, None)
        self.addCleanup(set_current_all_accessible, False)
        set_current_tenant(self.tenant_a)
        set_current_all_accessible(True)
        with patch("core.reports.orchestration.get_current_user", return_value=self.user):
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [], None)

    def test_ambient_tenant_scope_is_checked_without_arguments(self):
        provider = get_report_provider("asset_summary")
        self.addCleanup(set_current_tenant, None)
        set_current_tenant(self.tenant_a)
        with (
            patch("core.reports.orchestration.get_current_user", return_value=self.user),
            patch(
                "core.reports.orchestration._principal_covers_report_permissions",
                return_value=True,
            ) as check,
        ):
            _enforce_report_provider_permissions(provider, [], None)
        self.assertEqual([item.args[1] for item in check.call_args_list], [self.tenant_a])

    def test_aggregate_scope_resolution_failures_fail_closed(self):
        provider = get_report_provider("asset_summary")
        self.addCleanup(set_current_all_accessible, False)
        set_current_all_accessible(True)
        broken_model = SimpleNamespace(
            _base_manager=SimpleNamespace(filter=Mock(side_effect=RuntimeError("enumeration failed")))
        )
        with (
            patch("core.reports.orchestration.get_current_user", return_value=self.user),
            patch("core.reports.orchestration.tenant_model", return_value=broken_model),
        ):
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [], None)

        empty_model = SimpleNamespace(_base_manager=SimpleNamespace(filter=Mock(return_value=_ScopeQuery([]))))
        with (
            patch("core.reports.orchestration.get_current_user", return_value=self.user),
            patch("core.reports.orchestration.tenant_model", return_value=empty_model),
            patch("core.reports.orchestration.accessible_tenant_ids", return_value={1}),
        ):
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [], None)

        with (
            patch("core.reports.orchestration.get_current_user", return_value=self.user),
            patch(
                "core.reports.orchestration.get_current_scope_conflict",
                side_effect=RuntimeError("conflict probe failed"),
            ),
        ):
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [], None)

    def test_global_authorization_path_enumerates_live_tenants(self):
        """The truly global path keeps the live-tenant evaluation (defense in depth)."""
        tenant_query = Mock()
        tenant_query.values_list.return_value = (1, 2)
        tenant_model_value = SimpleNamespace(_base_manager=SimpleNamespace(filter=Mock(return_value=tenant_query)))
        permissions = ("assets.view_asset",)
        with (
            patch("core.reports.orchestration.tenant_model", return_value=tenant_model_value),
            patch(
                "core.reports.orchestration.build_accessible_tenant_permissions_map",
                return_value={1: (frozenset(permissions), None), 2: (frozenset(permissions), None)},
            ),
        ):
            authorized, live_ids = _report_compilation_authorized(
                user=self.user,
                principal_resolved=True,
                principal_active=True,
                principal_superuser=False,
                scope_tenants=(),
                permissions=permissions,
            )
        self.assertTrue(authorized)
        self.assertEqual(live_ids, (1, 2))

        with patch("core.reports.orchestration.tenant_model", side_effect=RuntimeError("enumeration failed")):
            self.assertEqual(_live_tenant_ids(), ())

    def test_empty_declaration_and_principal_resolver_error_deny(self):
        provider = Mock(report_type="empty_provider")
        provider.required_permissions.return_value = ()
        with patch("core.reports.orchestration.get_current_user", return_value=self.user):
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [], self.tenant_a)

        provider.required_permissions.return_value = ("assets.view_asset",)
        with patch("core.reports.orchestration.get_current_user", side_effect=RuntimeError("resolver failed")):
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [], self.tenant_a)

        provider.required_permissions.side_effect = RuntimeError("declaration failed")
        with patch("core.reports.orchestration.get_current_user", return_value=self.user):
            with self.assertRaises(ReportPermissionDenied):
                _enforce_report_provider_permissions(provider, [], self.tenant_a)

    def test_html_csv_xlsx_mail_and_download_headers_carry_disclosure(self):
        disclosure = "Showing <first> 2 of 6 rows & matching records."
        context_data = {"disclosure_text": disclosure}
        custom_html = SimpleNamespace(template_content="<p>Custom</p>")
        rendered = render_report_html(context_data, custom_html)
        self.assertIn('<aside class="report-output-disclosure" data-report-disclosure="true">', rendered)
        self.assertIn("&lt;first&gt;", rendered)
        self.assertNotIn("style=", rendered)

        polished = render_report_html(
            {
                "request": SimpleNamespace(csp_nonce="nonce"),
                "report_name": "Disclosure",
                "description": "",
                "generated_at": datetime(2026, 1, 1),
                "headers": ["Asset Tag"],
                "grouped_data": {},
                "summary_cards": [],
                "distribution_chart": "",
                "style_preset": "default",
                "is_compact": False,
                "is_financial": False,
                "is_sample": False,
                "row_limit": 500,
                "truncated": True,
                "total_rows": 600,
                "disclosure_text": "Showing the first 500 of 600 matching rows.",
            }
        )
        self.assertIn('<div class="disclosure">Showing the first 500 of 600 matching rows.</div>', polished)

        csv_text = render_report_csv(["Asset Tag"], [], disclosure_text=disclosure)
        self.assertEqual(list(csv.reader(io.StringIO(csv_text)))[-2:], [[], [disclosure]])

        from openpyxl import load_workbook

        workbook = load_workbook(
            io.BytesIO(report_xlsx_bytes(["Asset Tag"], [{"Asset Tag": "A-1"}], disclosure_text=disclosure))
        )
        worksheet = workbook.active
        self.assertIsNone(worksheet.cell(row=3, column=1).value)
        self.assertEqual(worksheet.cell(row=4, column=1).value, disclosure)
        self.assertTrue(worksheet.cell(row=4, column=1).font.italic)

        mail_body = _attachment_email_body("CSV", "Assets", disclosure_text=disclosure)
        self.assertTrue(mail_body.endswith(f"\n\n{disclosure}"))

        response = _apply_report_disclosure_headers(
            HttpResponse("csv"),
            {"truncated": True, "row_limit": 500, "total_rows": 600, "is_sample": False},
        )
        self.assertEqual(response["X-Report-Truncated"], "true")
        self.assertEqual(response["X-Report-Row-Window"], "500")
        self.assertEqual(response["X-Report-Total-Rows"], "600")
        sample_response = _apply_report_disclosure_headers(
            HttpResponse("html"), {"truncated": False, "is_sample": True}
        )
        self.assertEqual(sample_response["X-Report-Sample"], "true")

    def test_canonical_csv_uses_selected_columns_and_sanitizes_formula_cells(self):
        csv_text = render_report_csv(["Asset Tag"], [{"Asset Tag": "=SUM(A1)"}])

        self.assertEqual(list(csv.reader(io.StringIO(csv_text))), [["Asset Tag"], ["'=SUM(A1)"]])


class ReportColumnParityTests(SimpleTestCase):
    def test_every_provider_cell_key_is_offered_by_the_report_template_form(self):
        from assets.reports import (
            AssetDisposalEolReportProvider,
            AssetMaintenanceReportProvider,
            AssetSummaryReportProvider,
            WarrantyExpirationReportProvider,
        )
        from compliance.reports import CustodyComplianceReportProvider
        from licenses.reports import LicenseUtilizationReportProvider
        from software.reports import SoftwareInventoryReportProvider

        providers = (
            AssetSummaryReportProvider(),
            LicenseUtilizationReportProvider(),
            SubscriptionRenewalsReportProvider(),
            AssetMaintenanceReportProvider(),
            AssetDepreciationReportProvider(),
            SoftwareInventoryReportProvider(),
            ContractRenewalsReportProvider(),
            WarrantyExpirationReportProvider(),
            AssetDisposalEolReportProvider(),
            HardwareInventoryReportProvider(),
            CustodyComplianceReportProvider(),
        )
        offered = {key for key, _label in ReportTemplateForm.COLUMN_CHOICES}

        for provider in providers:
            missing = set(provider.cells) - offered
            self.assertFalse(
                missing, f"{provider.report_type} provider columns missing from the form: {sorted(missing)}"
            )
